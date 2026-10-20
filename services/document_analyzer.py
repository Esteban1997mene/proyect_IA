import io
import re
from datetime import datetime

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


# ============================================================
# INTI PROJECTS AI
# MOTOR DE ANÁLISIS DOCUMENTAL
# ============================================================


# ============================================================
# UTILIDADES
# ============================================================

def limpiar_texto(texto):
    if texto is None:
        return ""

    texto = str(texto)

    texto = texto.replace("\x00", " ")
    texto = re.sub(r"\r\n?", "\n", texto)
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)

    return texto.strip()


def limitar_texto(texto, limite=250000):
    """
    Evita almacenar textos excesivamente grandes
    en una sola variable de sesión.
    """

    texto = limpiar_texto(texto)

    if len(texto) <= limite:
        return texto

    return texto[:limite]


# ============================================================
# PDF
# ============================================================

def extraer_texto_pdf(bytes_archivo):

    try:
        archivo = io.BytesIO(bytes_archivo)

        lector = PdfReader(archivo)

        paginas = []

        for numero, pagina in enumerate(
            lector.pages,
            start=1
        ):

            try:
                texto = pagina.extract_text() or ""

            except Exception:
                texto = ""

            if texto.strip():

                paginas.append(
                    f"\n--- PÁGINA {numero} ---\n"
                    f"{texto}"
                )

        texto_final = "\n".join(paginas)

        return limpiar_texto(
            texto_final
        )

    except Exception as error:

        raise ValueError(
            f"No fue posible leer el PDF: {error}"
        )


# ============================================================
# DOCX
# ============================================================

def extraer_texto_docx(bytes_archivo):

    try:
        archivo = io.BytesIO(
            bytes_archivo
        )

        documento = Document(
            archivo
        )

        contenido = []

        # ----------------------------------------------------
        # PÁRRAFOS
        # ----------------------------------------------------

        for parrafo in documento.paragraphs:

            texto = parrafo.text.strip()

            if texto:
                contenido.append(
                    texto
                )

        # ----------------------------------------------------
        # TABLAS
        # ----------------------------------------------------

        for numero_tabla, tabla in enumerate(
            documento.tables,
            start=1
        ):

            contenido.append(
                f"\n--- TABLA {numero_tabla} ---"
            )

            for fila in tabla.rows:

                valores = []

                for celda in fila.cells:

                    valor = limpiar_texto(
                        celda.text
                    )

                    valores.append(
                        valor
                    )

                contenido.append(
                    " | ".join(
                        valores
                    )
                )

        return limpiar_texto(
            "\n".join(
                contenido
            )
        )

    except Exception as error:

        raise ValueError(
            f"No fue posible leer el DOCX: {error}"
        )


# ============================================================
# XLSX
# ============================================================

def extraer_texto_xlsx(bytes_archivo):

    try:
        archivo = io.BytesIO(
            bytes_archivo
        )

        libro = load_workbook(
            archivo,
            data_only=True,
            read_only=True
        )

        contenido = []

        for hoja in libro.worksheets:

            contenido.append(
                f"\n--- HOJA: {hoja.title} ---"
            )

            for fila in hoja.iter_rows(
                values_only=True
            ):

                valores = []

                for valor in fila:

                    if valor is None:
                        valores.append("")
                    else:
                        valores.append(
                            str(valor)
                        )

                if any(
                    valor.strip()
                    for valor in valores
                ):

                    contenido.append(
                        " | ".join(
                            valores
                        )
                    )

        return limpiar_texto(
            "\n".join(
                contenido
            )
        )

    except Exception as error:

        raise ValueError(
            f"No fue posible leer el XLSX: {error}"
        )


# ============================================================
# EXTRACCIÓN SEGÚN EXTENSIÓN
# ============================================================

def extraer_texto_documento(
    nombre,
    bytes_archivo
):

    nombre_lower = nombre.lower()

    if nombre_lower.endswith(".pdf"):

        return extraer_texto_pdf(
            bytes_archivo
        )

    if nombre_lower.endswith(".docx"):

        return extraer_texto_docx(
            bytes_archivo
        )

    if nombre_lower.endswith(".xlsx"):

        return extraer_texto_xlsx(
            bytes_archivo
        )

    raise ValueError(
        "Formato documental no soportado "
        f"para análisis: {nombre}"
    )


# ============================================================
# CLASIFICACIÓN DEL DOCUMENTO
# ============================================================

def clasificar_documento(
    nombre,
    texto
):

    contenido = (
        f"{nombre} {texto[:10000]}"
    ).lower()

    reglas = [
        (
            "PLIEGO_TDR",
            [
                "términos de referencia",
                "terminos de referencia",
                "pliego de condiciones",
                "solicitud de ofertas",
                "solicitud de propuesta"
            ]
        ),
        (
            "ESTUDIO_PREVIO",
            [
                "estudios previos",
                "estudio previo",
                "estudio de conveniencia"
            ]
        ),
        (
            "ANEXO_TECNICO",
            [
                "anexo técnico",
                "anexo tecnico",
                "especificaciones técnicas",
                "especificaciones tecnicas",
                "ficha técnica",
                "ficha tecnica"
            ]
        ),
        (
            "PRESUPUESTO",
            [
                "presupuesto oficial",
                "presupuesto",
                "apu",
                "análisis de precios unitarios",
                "analisis de precios unitarios"
            ]
        ),
        (
            "CRONOGRAMA",
            [
                "cronograma",
                "fecha de cierre",
                "calendario del proceso"
            ]
        ),
        (
            "FORMULARIO",
            [
                "formulario",
                "formato",
                "carta de presentación",
                "carta de presentacion"
            ]
        ),
        (
            "ADENDA",
            [
                "adenda",
                "modificación",
                "modificacion"
            ]
        )
    ]

    puntuaciones = {}

    for categoria, palabras in reglas:

        puntuacion = 0

        for palabra in palabras:

            if palabra in contenido:
                puntuacion += 1

        puntuaciones[
            categoria
        ] = puntuacion

    mejor_categoria = max(
        puntuaciones,
        key=puntuaciones.get
    )

    if puntuaciones[
        mejor_categoria
    ] == 0:

        return "OTRO"

    return mejor_categoria


# ============================================================
# BUSCAR FRAGMENTOS
# ============================================================

def buscar_fragmentos(
    texto,
    palabras_clave,
    max_resultados=8,
    contexto=450
):

    if not texto:
        return []

    texto_limpio = limpiar_texto(
        texto
    )

    texto_lower = texto_limpio.lower()

    resultados = []

    posiciones_usadas = []

    for palabra in palabras_clave:

        palabra_lower = palabra.lower()

        inicio_busqueda = 0

        while True:

            posicion = texto_lower.find(
                palabra_lower,
                inicio_busqueda
            )

            if posicion == -1:
                break

            # Evitar fragmentos prácticamente duplicados
            duplicado = any(
                abs(
                    posicion - posicion_anterior
                ) < 250
                for posicion_anterior
                in posiciones_usadas
            )

            if not duplicado:

                inicio = max(
                    0,
                    posicion - contexto
                )

                fin = min(
                    len(texto_limpio),
                    posicion
                    + len(palabra)
                    + contexto
                )

                fragmento = texto_limpio[
                    inicio:fin
                ].strip()

                resultados.append(
                    fragmento
                )

                posiciones_usadas.append(
                    posicion
                )

            if len(resultados) >= max_resultados:
                return resultados

            inicio_busqueda = (
                posicion
                + len(palabra_lower)
            )

    return resultados


# ============================================================
# CONVERTIR FRAGMENTOS A TEXTO
# ============================================================

def fragmentos_a_texto(
    fragmentos
):

    if not fragmentos:
        return ""

    bloques = []

    for indice, fragmento in enumerate(
        fragmentos,
        start=1
    ):

        bloques.append(
            f"[Fragmento {indice}]\n"
            f"{fragmento}"
        )

    return "\n\n".join(
        bloques
    )


# ============================================================
# EXTRAER SECCIONES RELEVANTES
# ============================================================

def extraer_secciones(
    texto
):

    configuracion = {

        "objeto": [
            "objeto del contrato",
            "objeto contractual",
            "objeto de la contratación",
            "objeto de la contratacion",
            "objeto"
        ],

        "presupuesto": [
            "presupuesto oficial",
            "presupuesto estimado",
            "valor estimado",
            "valor del contrato",
            "presupuesto"
        ],

        "plazo": [
            "plazo de ejecución",
            "plazo de ejecucion",
            "duración del contrato",
            "duracion del contrato",
            "plazo contractual"
        ],

        "experiencia": [
            "experiencia específica",
            "experiencia especifica",
            "experiencia del proponente",
            "experiencia mínima",
            "experiencia minima"
        ],

        "personal": [
            "personal mínimo",
            "personal minimo",
            "equipo mínimo",
            "equipo minimo",
            "personal profesional",
            "personal clave"
        ],

        "financieros": [
            "capacidad financiera",
            "índice de liquidez",
            "indice de liquidez",
            "nivel de endeudamiento",
            "razón de cobertura",
            "razon de cobertura",
            "capacidad organizacional"
        ],

        "especificaciones": [
            "especificaciones técnicas",
            "especificaciones tecnicas",
            "requisitos técnicos",
            "requisitos tecnicos",
            "características técnicas",
            "caracteristicas tecnicas"
        ],

        "factores_evaluacion": [
            "factores de evaluación",
            "factores de evaluacion",
            "criterios de evaluación",
            "criterios de evaluacion",
            "asignación de puntaje",
            "asignacion de puntaje",
            "puntaje"
        ],

        "garantias": [
            "garantías",
            "garantias",
            "garantía de seriedad",
            "garantia de seriedad",
            "garantía única",
            "garantia unica"
        ],

        "causales_rechazo": [
            "causales de rechazo",
            "causal de rechazo",
            "rechazo de la oferta",
            "rechazo de la propuesta"
        ],

        "entregables": [
            "entregables",
            "productos a entregar",
            "productos esperados",
            "informes"
        ],

        "forma_pago": [
            "forma de pago",
            "condiciones de pago",
            "pagos"
        ]
    }

    resultado = {}

    for seccion, palabras in configuracion.items():

        fragmentos = buscar_fragmentos(
            texto,
            palabras
        )

        resultado[
            seccion
        ] = fragmentos_a_texto(
            fragmentos
        )

    return resultado


# ============================================================
# COBERTURA DEL ANÁLISIS
# ============================================================

def calcular_cobertura(
    resultado
):

    campos_criticos = [
        "objeto",
        "presupuesto",
        "plazo",
        "experiencia",
        "personal",
        "financieros",
        "especificaciones",
        "factores_evaluacion",
        "garantias",
        "causales_rechazo"
    ]

    encontrados = 0

    detalle = {}

    for campo in campos_criticos:

        existe = bool(
            resultado.get(
                campo,
                ""
            ).strip()
        )

        detalle[
            campo
        ] = existe

        if existe:
            encontrados += 1

    porcentaje = round(
        encontrados
        / len(campos_criticos)
        * 100
    )

    return porcentaje, detalle


# ============================================================
# ANALIZAR UN DOCUMENTO
# ============================================================

def analizar_documento(
    nombre,
    bytes_archivo
):

    texto = extraer_texto_documento(
        nombre,
        bytes_archivo
    )

    texto = limitar_texto(
        texto
    )

    categoria = clasificar_documento(
        nombre,
        texto
    )

    secciones = extraer_secciones(
        texto
    )

    return {

        "nombre":
            nombre,

        "categoria":
            categoria,

        "texto":
            texto,

        "caracteres":
            len(texto),

        "secciones":
            secciones,

        "estado":
            "ANALIZADO",

        "fecha_analisis":
            datetime.now().isoformat()
    }


# ============================================================
# CONSOLIDAR RESULTADOS
# ============================================================

def consolidar_resultados(
    documentos_analizados
):

    campos = [
        "objeto",
        "presupuesto",
        "plazo",
        "experiencia",
        "personal",
        "financieros",
        "especificaciones",
        "factores_evaluacion",
        "garantias",
        "causales_rechazo",
        "entregables",
        "forma_pago"
    ]

    resultado = {}

    for campo in campos:

        bloques = []

        for documento in documentos_analizados:

            secciones = documento.get(
                "secciones",
                {}
            )

            contenido = secciones.get(
                campo,
                ""
            )

            if contenido:

                bloques.append(
                    f"DOCUMENTO: "
                    f"{documento.get('nombre')}\n\n"
                    f"{contenido}"
                )

        resultado[
            campo
        ] = "\n\n".join(
            bloques
        )

    porcentaje, cobertura = (
        calcular_cobertura(
            resultado
        )
    )

    return {
        "resultado":
            resultado,

        "porcentaje_analisis":
            porcentaje,

        "cobertura":
            cobertura
    }


# ============================================================
# ANALIZAR CONJUNTO DOCUMENTAL
# ============================================================

def analizar_documentos(
    documentos
):

    documentos_analizados = []

    errores = []

    for documento in documentos:

        nombre = documento.get(
            "nombre",
            ""
        )

        bytes_archivo = documento.get(
            "bytes"
        )

        if not bytes_archivo:

            errores.append({
                "documento":
                    nombre,

                "error":
                    "No existen bytes del archivo."
            })

            continue

        try:

            resultado = analizar_documento(
                nombre,
                bytes_archivo
            )

            # Conservar metadata original
            resultado[
                "tipo_mime"
            ] = documento.get(
                "tipo_mime"
            )

            resultado[
                "tamano"
            ] = documento.get(
                "tamano"
            )

            documentos_analizados.append(
                resultado
            )

        except Exception as error:

            errores.append({
                "documento":
                    nombre,

                "error":
                    str(error)
            })

    consolidado = consolidar_resultados(
        documentos_analizados
    )

    return {

        "documentos":
            documentos_analizados,

        "resultado":
            consolidado[
                "resultado"
            ],

        "porcentaje_analisis":
            consolidado[
                "porcentaje_analisis"
            ],

        "cobertura":
            consolidado[
                "cobertura"
            ],

        "errores":
            errores,

        "fecha_analisis":
            datetime.now().isoformat()
    }
